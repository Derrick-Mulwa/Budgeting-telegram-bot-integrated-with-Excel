import time
import asyncio
from datetime import datetime
import json
from telegram import *
from telegram.ext import *
from openpyxl import load_workbook

# Bot details
TOKEN = '6795785508:AAH5AGjkJQj30Elsl8PuC_YC6jN1pDkMF4g'
BOT_USERNAME = '@BajceticBot'
MY_CHAT_ID = '1486454053'


# create a list to mark progress
ongoingProcess = [False, 'ProcessID', 'Process Description', 'Message to send']

# create a list for new items to buy later
NewItem = ['ItemName', 0]

# Create a function that will load the Excel file and create variable names for all the tables and worksheets
def load_excel():
    # Load the excel file

    wb = load_workbook(filename="BajeticAutomated.xlsx")

    # Create variables for all the worksheets and all the tables in each worksheet
    October = wb['October']
    OctoberBudget = October.tables['OctoberBudget']
    OctoberTotals = October.tables['OctoberTotals']
    OctoberSummary = October.tables['OctoberSummary']
    OctoberUnbudgeted = October.tables['OctoberUnbudgeted']

    october = [October, OctoberBudget, OctoberTotals, OctoberSummary, OctoberUnbudgeted]


    November = wb['November']
    NovemberBudget = November.tables['NovemberBudget']
    NovemberTotals = November.tables['NovemberTotals']
    NovemberSummary = November.tables['NovemberSummary']
    NovemberUnbudgeted = November.tables['NovemberUnbudgeted']

    november = [November, NovemberBudget, NovemberTotals,  NovemberSummary, NovemberUnbudgeted]

    December = wb['December']
    DecemberBudget = December.tables['DecemberBudget']
    DecemberTotals = December.tables['DecemberTotals']
    DecemberSummary = December.tables['DecemberSummary']
    DecemberUnbudgeted = December.tables['DecemberUnbudgeted']

    december = [December, DecemberBudget, DecemberTotals, DecemberSummary, DecemberUnbudgeted]

    LongTerm = wb['LongTerm']
    LongtermProjects = LongTerm.tables['LongtermProjects']

    longterm = [LongTerm, LongtermProjects]

    Items = wb['Items']
    ItemsTable = Items.tables['Items']

    items = [Items, ItemsTable]

    return [october, november, december, longterm, items, wb]



async def command_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    load_excel()
    await update.message.reply_text('Hello')


async def command_newitem(update: Update, context: ContextTypes.DEFAULT_TYPE, ongoingProcess = ongoingProcess):
    # Check if there is an ongoing process. If not, get the data for the new item and add to the excel db
    message_id = update.message.message_id
    print('Intiated Add New Item')
    if ongoingProcess[0] is False:
        ongoingProcess[0] = True
        ongoingProcess[1] = 'AddNewItem:Step1'
        ongoingProcess[2] = 'Adding a new Item'
        ongoingProcess[3] = 'You have an ongoing process: Adding a new Item. \n\nSend "Cancel Ongoing Process" to terminate the process'

        print('Requesting new name')
        await update.message.reply_text('Enter the name of the new item', reply_to_message_id=message_id)

    else:
        await update.message.reply_text(f'{ongoingProcess[3]}', reply_to_message_id=message_id)


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE, ongoingProcess = ongoingProcess, NewItem = NewItem):
    print(ongoingProcess)

    message_text = update.message.text
    message_id = update.message.message_id

    # Cancel any ongoing process and reset it using the "Cancel Ongoing Process" keyword
    if message_text.lower() == 'cancel ongoing process':
        print('Requested cancelling of ongoing process')
        ongoingProcess[0] = False
        ongoingProcess[1] = 'ProcessID'
        ongoingProcess[2] = 'Process Description'
        ongoingProcess[3] = 'Message to send'

        NewItem[0] = 'ItemName'
        NewItem[1] = 0

        await update.message.reply_text(f'Any ongoing process has been reset successfully', reply_to_message_id=message_id)


    elif (ongoingProcess[0] is True) and ongoingProcess[1] == 'AddNewItem:Step1':
        item_name = message_text
        if item_name != '':
            ongoingProcess[1] = 'AddNewItem:Step2'
            ongoingProcess[3] = f'You have an ongoing process: Adding price for new item: {item_name}.\n\nSend "Cancel Ongoing Process" to terminate the process'
            NewItem[0] = item_name
            print('Requesting price for new Item')
            await update.message.reply_text(f'Enter the price for the new item: {item_name}', reply_to_message_id=message_id)

        else:
            await update.message.reply_text(f'Message cannot be empty. Please re-enter the new item\'s name.', reply_to_message_id=message_id)



    # Adding a new item after collecting all the details
    elif (ongoingProcess[0] is True) and ongoingProcess[1] == 'AddNewItem:Step2':
        item_price = message_text
        try:
            item_price = int(item_price)
            if item_price <= 0:
                await update.message.reply_text(text="Invalid Price! Please enter a number greater than 0",
                                                reply_to_message_id=message_id)

            else:
                NewItem[1] = item_price

                # Write on Excel
                # print(NewItem)

                # Load the Excel
                excel = load_excel()

                print('Loaded Excel')
                Items_worksheet = excel[4][0]
                Items_table = excel[4][1]

                workbook = excel[5]

                print('Initialized Item table')

                ref = Items_table.ref

                print(f'Ref: {ref}')
                last_row_number = ref[ref.index(":") + 1:][[index for index, char in enumerate(ref[ref.index(":") + 1:]) if char.isdigit()][0]:]
                new_row_number = int(last_row_number)+1
                print(f'last_row_number: {last_row_number}')
                print(f'New item: {NewItem}')

                Items_worksheet[f'A{new_row_number}'] = NewItem[0]
                Items_worksheet[f'B{new_row_number}'] = NewItem[1]

                table_start = ref[:ref.index(":") + 1]
                print(f'Table start: {table_start}')

                new_table_end_column = ref[ref.index(":") + 1:][:[index for index, char in
                                                                  enumerate(ref[ref.index(":") + 1:]) if
                                                                  char.isdigit()][0]]
                new_table_end_row = new_row_number
                table_end = f'{new_table_end_column}{new_table_end_row}'

                new_table_ref = f'{table_start}{table_end}'

                Items_table.ref = new_table_ref

                workbook.save("BajeticAutomated.xlsx")

                await update.message.reply_text(text=f"New item added successfully!\nName: {NewItem[0]}\nPrice: {NewItem[1]}",
                                                reply_to_message_id=message_id)

                print('Resetting status')
                ongoingProcess[0] = False
                ongoingProcess[1] = 'ProcessID'
                ongoingProcess[2] = 'Process Description'
                ongoingProcess[3] = 'Message to send'

                NewItem[0] = 'ItemName'
                NewItem[1] = 0
















        except:
            await update.message.reply_text(text="Invalid Price! Please enter a number greater than 0", reply_to_message_id=message_id)





async def command_cancelongoingprocess(update: Update, context: ContextTypes.DEFAULT_TYPE, ongoingProcess = ongoingProcess, NewItem = NewItem):
    print(ongoingProcess)

    message_id = update.message.message_id

    ongoingProcess[0] = False
    ongoingProcess[1] = 'ProcessID'
    ongoingProcess[2] = 'Process Description'
    ongoingProcess[3] = 'Message to send'

    NewItem[0] = 'ItemName'
    NewItem[1] = 0

    await update.message.reply_text(f'Any ongoing process has been reset successfully', reply_to_message_id=message_id)





if __name__ == '__main__':
    print('Starting bot')
    app = Application.builder().token(TOKEN).build()

    # add command handlers for bot commands
    app.add_handler(CommandHandler('start', command_start))
    app.add_handler(CommandHandler('newitem', command_newitem))
    app.add_handler(CommandHandler('cancelongoingprocess', command_cancelongoingprocess))
    app.add_handler(CommandHandler('newproject', command_newproject))



    # Querry handler
    # app.add_handler(CallbackQueryHandler(callbackhandler))

    # message Handler
    app.add_handler(MessageHandler(filters.TEXT, handle_message))

    # Add error handler
    app.add_error_handler(error)
    print('Polling...')
    app.run_polling(poll_interval=1)